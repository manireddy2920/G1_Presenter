"""Optional simple test deck; install python-pptx before running."""
from pptx import Presentation
r=Presentation()
for title,body,narration in [
 ('Welcome to the engineering lab','Your humanoid presentation host','Welcome. I will demonstrate visitor detection, display selection and a narrated presentation.'),
 ('How this presentation works','Camera detects a visitor\nA configured display is selected\nPowerPoint starts and speech follows each slide','The camera detects a visitor. I select an available configured display, open PowerPoint, and narrate the slides.'),
 ('Robot and display working together','PowerPoint on a Windows display PC\nAudio through the TV speakers\nG1 gestures through SDK2','The display computer runs PowerPoint and sends speech to the television speakers. The robot supplies a configured greeting gesture. Thank you for visiting.')]:
 s=r.slides.add_slide(r.slide_layouts[1]);s.shapes.title.text=title;s.placeholders[1].text=body;s.notes_slide.notes_text_frame.text=narration
r.save('demo.pptx')
