"""Export slide narration as offline WAV files on Windows."""
import argparse
from pathlib import Path
from pptx import Presentation
import win32com.client
p=argparse.ArgumentParser();p.add_argument('pptx');p.add_argument('--out',default='narration_audio');a=p.parse_args()
out=Path(a.out).resolve();out.mkdir(parents=True,exist_ok=True)
voice=win32com.client.Dispatch('SAPI.SpVoice')
for n,s in enumerate(Presentation(a.pptx).slides,1):
 text=s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ''
 if not text: text=' '.join(shape.text for shape in s.shapes if shape.has_text_frame)
 stream=win32com.client.Dispatch('SAPI.SpFileStream')
 stream.Open(str(out/f'slide_{n:03}.wav'),3)
 try:
  voice.AudioOutputStream=stream
  voice.Speak(text or f'Slide {n}',0)
 finally: stream.Close()
print('WAV narration saved to',out)
